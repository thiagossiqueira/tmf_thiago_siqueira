#!/usr/bin/env python
"""
Build a small Bloomberg/YAS validation workbook for the corporate quote matrix.

Why this script exists
----------------------
The legacy ``ya.v1.xlsx`` file is not homogeneous: some source-specific
``PX_LAST`` series appear to be price/PU quotes, while others appear to be
already quoted as yields. Before converting the complete 103-bond thesis
universe, this script creates a controlled sample that can be refreshed in
Bloomberg Excel using a historical YAS price override.

The script does not modify any source workbook and does not calculate final
corporate yields locally. It creates:

- ``validation_sample``: a small, diverse sample with live Bloomberg formulas;
- ``candidate_universe``: all eligible 103-bond monthly observations, without
  Bloomberg formulas;
- audit and instruction sheets.

Expected counts under the audited thesis snapshot
-------------------------------------------------
- 103 thesis bonds
- 4,651 eligible bond-month quotes
- 4,266 rows with an old IPCA spread-output reference
- 4,229 rows present in the submitted econometric panel
- 385 eligible rows without an old spread-output reference
"""

from __future__ import annotations

import argparse
import math
import sys
from pathlib import Path
from typing import Iterable

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


REPO_ROOT = Path(__file__).resolve().parents[1]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from src.config import CONFIG  # noqa: E402
from src.core.curve_builder import load_real_curve_support  # noqa: E402
from src.core.windowing import build_observation_windows  # noqa: E402
from src.utils.file_io import (  # noqa: E402
    load_corp_bond_data,
    load_ipca_surface,
    load_yield_surface,
)


DEFAULT_PANEL_PATH = (
    REPO_ROOT
    / "datos_y_modelos"
    / "db"
    / "econometrics_db"
    / "panel_main_and_robustness.xlsx"
)

DEFAULT_OLD_SPREAD_PATH = (
    REPO_ROOT
    / "datos_y_modelos"
    / "db"
    / "output_calculated_spreads"
    / "output_spreads_di_and_ipca.xlsx"
)

DEFAULT_OUTPUT_PATH = (
    REPO_ROOT
    / "datos_y_modelos"
    / "db"
    / "brazil_domestic_corp_bonds"
    / "corp_yas_validation_sample.v1.xlsx"
)

EXPECTED_BONDS = 103
EXPECTED_ELIGIBLE_ROWS = 4_651
EXPECTED_OLD_OUTPUT_ROWS = 4_266
EXPECTED_OLD_PANEL_ROWS = 4_229
EXPECTED_NO_OLD_OUTPUT_ROWS = 385

ANCHOR_CASES = (
    ("EJ422135 Corp", "2012-12-28"),  # clearly price/PU-like
    ("AM250786 Corp", "2017-06-30"),  # stored quote matches old YTM
)


class ValidationError(RuntimeError):
    """Raised when the audit universe violates an expected invariant."""


def _clean_id(value: object) -> str:
    return str(value).strip()


def _require_columns(
    df: pd.DataFrame,
    required: Iterable[str],
    dataset_name: str,
) -> None:
    missing = [column for column in required if column not in df.columns]
    if missing:
        raise KeyError(
            f"{dataset_name} is missing required columns: {missing}"
        )


def _month_key(values: pd.Series) -> pd.Series:
    return pd.to_datetime(values, errors="coerce").dt.to_period("M")


def _safe_numeric(value: object) -> float:
    converted = pd.to_numeric(value, errors="coerce")
    return float(converted) if pd.notna(converted) else math.nan


def _classify_quote(row: pd.Series) -> str:
    """Create a preliminary, non-final quote classification.

    Small negative values are retained: the audit found continuous negative
    series and an exact match to a legacy corporate YTM. Exact zero and the
    repeated -100 value are retained for audit but flagged separately.
    """
    stored_quote = _safe_numeric(row["stored_quote"])
    old_yield = _safe_numeric(row["old_corporate_yield_pct"])

    if not np.isfinite(stored_quote):
        return "INVALID_MISSING"

    if np.isfinite(old_yield) and abs(stored_quote - old_yield) <= 0.01:
        return "STORED_QUOTE_MATCHES_OLD_YIELD"

    if stored_quote == -100.0:
        return "INVALID_SENTINEL_MINUS_100"
    if stored_quote == 0.0:
        return "ZERO_QUOTE_REVIEW"

    if np.isfinite(old_yield):
        if abs(stored_quote) >= 100.0:
            return "PRICE_LIKE_WITH_OLD_YIELD_REFERENCE"
        return "QUOTE_DIFFERS_FROM_OLD_YIELD"

    if stored_quote < 0.0:
        return "NEGATIVE_YIELD_LIKE_NO_OLD_REFERENCE"
    if abs(stored_quote) >= 100.0:
        return "PRICE_LIKE_NO_OLD_REFERENCE"
    if abs(stored_quote) <= 20.0:
        return "LOW_VALUE_NO_OLD_REFERENCE"
    return "MID_RANGE_NO_OLD_REFERENCE"


def _one_representative_per_bond(df: pd.DataFrame) -> pd.DataFrame:
    """Choose a deterministic middle observation for each bond."""
    representatives: list[pd.Series] = []
    for _, group in df.sort_values(["bond_id", "obs_date"]).groupby(
        "bond_id",
        sort=True,
    ):
        representatives.append(group.iloc[len(group) // 2])
    if not representatives:
        return df.iloc[0:0].copy()
    return pd.DataFrame(representatives).reset_index(drop=True)


def _evenly_spaced_rows(df: pd.DataFrame, n: int) -> pd.DataFrame:
    if n <= 0 or df.empty:
        return df.iloc[0:0].copy()
    if len(df) <= n:
        return df.copy()
    positions = np.linspace(0, len(df) - 1, n).round().astype(int)
    return df.iloc[np.unique(positions)].copy()


def _select_diverse_sample(
    candidates: pd.DataFrame,
    per_group: int,
) -> pd.DataFrame:
    """Select anchors plus diverse observations from four audit groups."""
    selected_parts: list[pd.DataFrame] = []
    selected_keys: set[tuple[str, pd.Timestamp]] = set()

    def add_rows(rows: pd.DataFrame, sample_group: str) -> None:
        if rows.empty:
            return
        rows = rows.copy()
        rows["sample_group"] = sample_group
        rows = rows[
            ~rows.apply(
                lambda r: (r["bond_id"], r["obs_date"]) in selected_keys,
                axis=1,
            )
        ]
        if rows.empty:
            return
        for _, row in rows.iterrows():
            selected_keys.add((row["bond_id"], row["obs_date"]))
        selected_parts.append(rows)

    # Explicit anchor cases discussed during the audit.
    anchor_rows: list[pd.DataFrame] = []
    for bond_id, date_text in ANCHOR_CASES:
        obs_date = pd.Timestamp(date_text)
        match = candidates[
            (candidates["bond_id"] == bond_id)
            & (candidates["obs_date"] == obs_date)
        ]
        if not match.empty:
            anchor_rows.append(match.head(1))
    if anchor_rows:
        add_rows(pd.concat(anchor_rows, ignore_index=True), "ANCHOR")

    group_specs = [
        (
            "PRICE_LIKE",
            candidates[
                candidates["preliminary_classification"].isin(
                    [
                        "PRICE_LIKE_WITH_OLD_YIELD_REFERENCE",
                        "PRICE_LIKE_NO_OLD_REFERENCE",
                    ]
                )
            ],
        ),
        (
            "STORED_MATCHES_OLD_YIELD",
            candidates[
                candidates["preliminary_classification"]
                == "STORED_QUOTE_MATCHES_OLD_YIELD"
            ],
        ),
        (
            "DIFFERS_FROM_OLD_YIELD",
            candidates[
                candidates["preliminary_classification"]
                == "QUOTE_DIFFERS_FROM_OLD_YIELD"
            ],
        ),
        (
            "NEGATIVE_YIELD_LIKE",
            candidates[
                candidates["preliminary_classification"]
                == "NEGATIVE_YIELD_LIKE_NO_OLD_REFERENCE"
            ],
        ),
        (
            "INVALID_OR_ZERO_REVIEW",
            candidates[
                candidates["preliminary_classification"].isin(
                    ["INVALID_SENTINEL_MINUS_100", "ZERO_QUOTE_REVIEW"]
                )
            ],
        ),
        (
            "NO_OLD_REFERENCE",
            candidates[candidates["old_corporate_yield_pct"].isna()],
        ),
    ]

    for label, group in group_specs:
        if group.empty:
            continue

        available = group[
            ~group.apply(
                lambda r: (r["bond_id"], r["obs_date"]) in selected_keys,
                axis=1,
            )
        ]

        if available.empty:
            continue

        representatives = _one_representative_per_bond(available)
        chosen = _evenly_spaced_rows(
            representatives.sort_values(
                ["stored_quote", "bond_id", "obs_date"]
            ),
            per_group,
        )
        add_rows(chosen, label)

    if not selected_parts:
        raise ValidationError("No validation sample rows were selected.")

    sample = pd.concat(selected_parts, ignore_index=True)
    return sample.sort_values(
        ["sample_group", "bond_id", "obs_date"]
    ).reset_index(drop=True)


def _style_workbook(path: Path) -> None:
    """Apply light audit-friendly formatting and Bloomberg formulas."""
    workbook = load_workbook(path)

    try:
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.calculation.calcMode = "auto"
    except AttributeError:
        # Older openpyxl versions may expose calculation settings differently.
        pass

    dark_fill = PatternFill("solid", fgColor="1F4E78")
    light_blue_fill = PatternFill("solid", fgColor="D9EAF7")
    yellow_fill = PatternFill("solid", fgColor="FFF2CC")
    orange_fill = PatternFill("solid", fgColor="FCE4D6")
    white_font = Font(color="FFFFFF", bold=True)

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False

        for cell in sheet[1]:
            cell.fill = dark_fill
            cell.font = white_font
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        for column_cells in sheet.columns:
            letter = get_column_letter(column_cells[0].column)
            max_length = 0
            for cell in column_cells[:300]:
                value = "" if cell.value is None else str(cell.value)
                max_length = max(max_length, len(value))
            sheet.column_dimensions[letter].width = min(
                max(max_length + 2, 11),
                48,
            )

    sample_sheet = workbook["validation_sample"]
    headers = {
        cell.value: cell.column
        for cell in sample_sheet[1]
    }

    required_headers = {
        "bond_id",
        "obs_date",
        "settle_date",
        "stored_quote",
        "yas_request_eligible",
        "old_corporate_yield_pct",
        "yas_yield_if_treated_as_price",
        "yas_minus_old_yield_bp",
        "automated_check",
        "analyst_review",
    }
    missing = sorted(required_headers - set(headers))
    if missing:
        raise ValidationError(
            f"Validation sheet missing formula columns: {missing}"
        )

    bond_col = get_column_letter(headers["bond_id"])
    settle_col = get_column_letter(headers["settle_date"])
    stored_col = get_column_letter(headers["stored_quote"])
    eligible_col = get_column_letter(headers["yas_request_eligible"])
    old_yield_col = get_column_letter(headers["old_corporate_yield_pct"])
    yas_col = get_column_letter(headers["yas_yield_if_treated_as_price"])
    yas_diff_col = get_column_letter(headers["yas_minus_old_yield_bp"])
    check_col = get_column_letter(headers["automated_check"])
    review_col = get_column_letter(headers["analyst_review"])

    for row_number in range(2, sample_sheet.max_row + 1):
        # Formula intentionally mirrors the historical Excel/YAS method.
        sample_sheet[f"{yas_col}{row_number}"] = (
            f'=IF(OR({eligible_col}{row_number}<>TRUE,'
            f'{bond_col}{row_number}="",'
            f'{stored_col}{row_number}="",'
            f'{settle_col}{row_number}=""),"",'
            f'BDP({bond_col}{row_number},"YAS_BOND_YLD",'
            f'"YAS_BOND_PX",{stored_col}{row_number},'
            f'"SETTLE_DT",TEXT({settle_col}{row_number},"YYYYMMDD")))'
        )
        sample_sheet[f"{yas_diff_col}{row_number}"] = (
            f'=IF(OR(NOT(ISNUMBER({yas_col}{row_number})),'
            f'NOT(ISNUMBER({old_yield_col}{row_number}))),"",'
            f'100*({yas_col}{row_number}-{old_yield_col}{row_number}))'
        )
        sample_sheet[f"{check_col}{row_number}"] = (
            f'=IF(NOT(ISNUMBER({yas_col}{row_number})),"PENDING/ERROR",'
            f'IF(AND(ISNUMBER({old_yield_col}{row_number}),'
            f'ABS(100*({yas_col}{row_number}-{old_yield_col}{row_number}))<=5),'
            f'"PRICE->YAS MATCHES OLD YIELD",'
            f'IF(AND(ISNUMBER({old_yield_col}{row_number}),'
            f'ABS(100*({stored_col}{row_number}-{old_yield_col}{row_number}))<=5),'
            f'"STORED QUOTE MATCHES OLD YIELD","REVIEW")))'
        )

        sample_sheet[f"{yas_col}{row_number}"].fill = yellow_fill
        sample_sheet[f"{yas_diff_col}{row_number}"].fill = light_blue_fill
        sample_sheet[f"{check_col}{row_number}"].fill = light_blue_fill
        sample_sheet[f"{review_col}{row_number}"].fill = orange_fill

    for sheet_name in ["validation_sample", "candidate_universe"]:
        sheet = workbook[sheet_name]
        header_map = {cell.value: cell.column for cell in sheet[1]}
        for date_header in ["obs_date", "settle_date"]:
            if date_header in header_map:
                letter = get_column_letter(header_map[date_header])
                for row_number in range(2, sheet.max_row + 1):
                    sheet[f"{letter}{row_number}"].number_format = "yyyy-mm-dd"
        for numeric_header in [
            "stored_quote",
            "old_corporate_yield_pct",
            "old_benchmark_yield_pct",
            "old_spread_field",
            "stored_minus_old_yield_bp",
            "yas_yield_if_treated_as_price",
            "yas_minus_old_yield_bp",
        ]:
            if numeric_header in header_map:
                letter = get_column_letter(header_map[numeric_header])
                for row_number in range(2, sheet.max_row + 1):
                    sheet[f"{letter}{row_number}"].number_format = "0.000000"

    workbook.save(path)


def build_validation_workbook(
    panel_path: Path,
    old_spread_path: Path,
    output_path: Path,
    per_group: int = 8,
    strict: bool = True,
) -> Path:
    if not panel_path.exists():
        raise FileNotFoundError(f"Panel not found: {panel_path}")
    if not old_spread_path.exists():
        raise FileNotFoundError(
            f"Old spread output not found: {old_spread_path}"
        )

    # --------------------------------------------------------------
    # 1. Fixed thesis bond universe and submitted panel reference
    # --------------------------------------------------------------
    panel = pd.read_excel(panel_path, sheet_name="ts")
    _require_columns(panel, ["bond_id", "obs_date"], "Econometric panel")
    panel = panel.copy()
    panel["bond_id"] = panel["bond_id"].map(_clean_id)
    panel["old_panel_obs_date"] = pd.to_datetime(
        panel["obs_date"],
        errors="coerce",
    )
    panel["month_key"] = _month_key(panel["old_panel_obs_date"])

    if panel["old_panel_obs_date"].isna().any():
        raise ValidationError("Panel contains invalid observation dates.")

    thesis_bond_ids = sorted(panel["bond_id"].unique().tolist())
    old_panel_pairs = set(zip(panel["bond_id"], panel["month_key"]))

    # --------------------------------------------------------------
    # 2. Metadata and legacy source-specific quote matrix
    # --------------------------------------------------------------
    corp = load_corp_bond_data(CONFIG["CORP_PATH"]).copy()
    _require_columns(corp, ["id", "MATURITY"], "Corporate metadata")
    corp["id"] = corp["id"].map(_clean_id)
    corp["MATURITY"] = pd.to_datetime(corp["MATURITY"], errors="coerce")
    corp = corp[corp["id"].isin(thesis_bond_ids)].copy()
    corp = corp.drop_duplicates("id", keep="first")

    missing_metadata = sorted(set(thesis_bond_ids) - set(corp["id"]))
    if missing_metadata:
        raise ValidationError(
            "Thesis bonds missing from metadata: "
            f"{missing_metadata[:20]}"
        )

    quotes = load_yield_surface(CONFIG["YA_PATH"]).copy()
    quotes.columns = [_clean_id(column) for column in quotes.columns]
    quotes.index = pd.to_datetime(quotes.index, errors="coerce")
    quotes = quotes[~quotes.index.isna()].sort_index()

    missing_quote_columns = sorted(
        set(thesis_bond_ids) - set(quotes.columns)
    )
    if missing_quote_columns:
        raise ValidationError(
            "Thesis bonds missing from legacy quote matrix: "
            f"{missing_quote_columns[:20]}"
        )

    observation_windows = build_observation_windows(
        corp,
        quotes,
        CONFIG["OBS_WINDOW"],
    )

    # --------------------------------------------------------------
    # 3. Valid monthly benchmark dates
    # --------------------------------------------------------------
    _, ntnb_yields = load_real_curve_support()
    wla_surface = load_ipca_surface(CONFIG["WLA_CURVE_PATH"]).copy()
    _require_columns(wla_surface, ["obs_date"], "WLA surface")
    wla_dates = pd.DatetimeIndex(
        pd.to_datetime(
            wla_surface["obs_date"],
            errors="coerce",
        ).dropna().unique()
    ).sort_values()

    valid_dates = (
        quotes.index
        .intersection(ntnb_yields.index)
        .intersection(wla_dates)
        .sort_values()
    )

    # --------------------------------------------------------------
    # 4. All eligible 103-bond monthly observations
    # --------------------------------------------------------------
    records: list[dict[str, object]] = []
    for _, bond in corp.sort_values("id").iterrows():
        bond_id = bond["id"]
        obs_start, obs_end = observation_windows[bond_id]
        eligible_dates = valid_dates[
            (valid_dates >= obs_start)
            & (valid_dates <= obs_end)
        ]
        selected_quotes = pd.to_numeric(
            quotes.loc[eligible_dates, bond_id],
            errors="coerce",
        ).dropna()

        for obs_date, stored_quote in selected_quotes.items():
            obs_date = pd.Timestamp(obs_date)
            month_key = obs_date.to_period("M")
            records.append(
                {
                    "request_id": (
                        f"{bond_id.replace(' ', '_')}_{obs_date:%Y%m%d}"
                    ),
                    "bond_id": bond_id,
                    "obs_date": obs_date,
                    "settle_date": obs_date,
                    "month_key": str(month_key),
                    "stored_quote": float(stored_quote),
                    "in_old_panel": (
                        bond_id,
                        month_key,
                    ) in old_panel_pairs,
                }
            )

    candidates = pd.DataFrame.from_records(records)
    if candidates.empty:
        raise ValidationError("No eligible candidate observations found.")

    candidates = candidates.sort_values(
        ["bond_id", "obs_date"]
    ).reset_index(drop=True)

    # --------------------------------------------------------------
    # 5. Old IPCA spread-output reference
    # --------------------------------------------------------------
    old_output = pd.read_excel(
        old_spread_path,
        sheet_name="values_only",
    )
    required_old_columns = [
        "Benchmark",
        "Bond ID",
        "Obs Date",
        "Corp Yield (%)",
        "Yield (%)",
        "Spread (bp)",
    ]
    _require_columns(
        old_output,
        required_old_columns,
        "Old spread output",
    )

    old_output = old_output[
        old_output["Benchmark"].astype(str).str.upper().eq("IPCA")
    ].copy()
    old_output["bond_id"] = old_output["Bond ID"].map(_clean_id)
    old_output["obs_date"] = pd.to_datetime(
        old_output["Obs Date"],
        errors="coerce",
    )
    old_output["month_key"] = (
        old_output["obs_date"]
        .dt.to_period("M")
        .astype(str)
    )
    old_output = old_output[
        old_output["bond_id"].isin(thesis_bond_ids)
    ].copy()

    old_reference = old_output[
        [
            "bond_id",
            "obs_date",
            "month_key",
            "Corp Yield (%)",
            "Yield (%)",
            "Spread (bp)",
        ]
    ].rename(
        columns={
            "obs_date": "old_output_obs_date",
            "Corp Yield (%)": "old_corporate_yield_pct",
            "Yield (%)": "old_benchmark_yield_pct",
            "Spread (bp)": "old_spread_field",
        }
    )

    duplicate_old_pairs = int(
        old_reference.duplicated(["bond_id", "month_key"]).sum()
    )

    if duplicate_old_pairs:
        raise ValidationError(
            "Old spread output contains duplicate bond-month pairs: "
            f"{duplicate_old_pairs}"
        )

    candidates = candidates.merge(
        old_reference,
        on=["bond_id", "month_key"],
        how="left",
        validate="one_to_one",
    )

    candidates["stored_minus_old_yield_bp"] = (
        100.0
        * (
            candidates["stored_quote"]
            - candidates["old_corporate_yield_pct"]
        )
    )
    candidates["preliminary_classification"] = candidates.apply(
        _classify_quote,
        axis=1,
    )

    # --------------------------------------------------------------
    # 6. Validation and deterministic sample selection
    # --------------------------------------------------------------
    unique_bonds = int(candidates["bond_id"].nunique())
    eligible_rows = int(len(candidates))
    old_output_rows = int(
        candidates["old_corporate_yield_pct"].notna().sum()
    )
    old_panel_rows = int(candidates["in_old_panel"].sum())
    no_old_output_rows = int(
        candidates["old_corporate_yield_pct"].isna().sum()
    )
    duplicate_requests = int(
        candidates["request_id"].duplicated().sum()
    )
    duplicate_bond_months = int(
        candidates.duplicated(["bond_id", "month_key"]).sum()
    )

    if duplicate_requests or duplicate_bond_months:
        raise ValidationError(
            "Duplicate observations detected: "
            f"request_id={duplicate_requests}, "
            f"bond_month={duplicate_bond_months}"
        )

    # Only positive, not-already-identified-as-yield quotes are candidates for
    # a Bloomberg YAS price override. Negative values stay in the audit.
    candidates["yas_request_eligible"] = (
        candidates["stored_quote"] > 0
    ) & (
        candidates["preliminary_classification"]
        != "STORED_QUOTE_MATCHES_OLD_YIELD"
    )

    actual = {
        "bonds": unique_bonds,
        "eligible_rows": eligible_rows,
        "old_output_rows": old_output_rows,
        "old_panel_rows": old_panel_rows,
        "no_old_output_rows": no_old_output_rows,
    }
    expected = {
        "bonds": EXPECTED_BONDS,
        "eligible_rows": EXPECTED_ELIGIBLE_ROWS,
        "old_output_rows": EXPECTED_OLD_OUTPUT_ROWS,
        "old_panel_rows": EXPECTED_OLD_PANEL_ROWS,
        "no_old_output_rows": EXPECTED_NO_OLD_OUTPUT_ROWS,
    }
    mismatches = {
        key: (actual[key], expected[key])
        for key in expected
        if actual[key] != expected[key]
    }
    if strict and mismatches:
        raise ValidationError(
            "Counts differ from the audited thesis snapshot: "
            f"{mismatches}. Investigate before using --no-strict."
        )

    sample = _select_diverse_sample(candidates, per_group=per_group)
    sample["yas_yield_if_treated_as_price"] = ""
    sample["yas_minus_old_yield_bp"] = ""
    sample["automated_check"] = ""
    sample["analyst_review"] = ""
    sample["notes"] = ""

    sample_columns = [
        "sample_group",
        "preliminary_classification",
        "request_id",
        "bond_id",
        "obs_date",
        "settle_date",
        "stored_quote",
        "yas_request_eligible",
        "old_corporate_yield_pct",
        "old_benchmark_yield_pct",
        "old_spread_field",
        "in_old_panel",
        "stored_minus_old_yield_bp",
        "yas_yield_if_treated_as_price",
        "yas_minus_old_yield_bp",
        "automated_check",
        "analyst_review",
        "notes",
    ]
    sample = sample[sample_columns]

    candidate_columns = [
        "request_id",
        "bond_id",
        "obs_date",
        "settle_date",
        "month_key",
        "stored_quote",
        "yas_request_eligible",
        "old_corporate_yield_pct",
        "old_benchmark_yield_pct",
        "old_spread_field",
        "in_old_panel",
        "stored_minus_old_yield_bp",
        "preliminary_classification",
    ]
    candidate_universe = candidates[candidate_columns].copy()

    classification_summary = (
        candidates.groupby("preliminary_classification", dropna=False)
        .agg(
            observations=("request_id", "size"),
            unique_bonds=("bond_id", "nunique"),
            old_output_references=(
                "old_corporate_yield_pct",
                lambda s: int(s.notna().sum()),
            ),
            old_panel_rows=("in_old_panel", "sum"),
            min_stored_quote=("stored_quote", "min"),
            median_stored_quote=("stored_quote", "median"),
            max_stored_quote=("stored_quote", "max"),
        )
        .reset_index()
        .sort_values("preliminary_classification")
    )

    bond_summary = (
        candidates.groupby("bond_id")
        .agg(
            eligible_rows=("request_id", "size"),
            old_output_rows=(
                "old_corporate_yield_pct",
                lambda s: int(s.notna().sum()),
            ),
            old_panel_rows=("in_old_panel", "sum"),
            first_obs_date=("obs_date", "min"),
            last_obs_date=("obs_date", "max"),
            min_stored_quote=("stored_quote", "min"),
            median_stored_quote=("stored_quote", "median"),
            max_stored_quote=("stored_quote", "max"),
        )
        .reset_index()
        .sort_values("bond_id")
    )

    summary = pd.DataFrame(
        [
            ("thesis_bonds", unique_bonds),
            ("eligible_bond_month_quotes", eligible_rows),
            ("old_ipca_output_references", old_output_rows),
            ("submitted_panel_rows", old_panel_rows),
            ("eligible_rows_without_old_output", no_old_output_rows),
            ("validation_sample_rows", len(sample)),
            ("valid_benchmark_dates", len(valid_dates)),
            ("duplicate_request_ids", duplicate_requests),
            ("duplicate_bond_months", duplicate_bond_months),
            ("zero_quote_rows", int((candidates["stored_quote"] == 0).sum())),
            ("minus_100_sentinel_rows", int((candidates["stored_quote"] == -100).sum())),
            (
                "negative_non_sentinel_rows",
                int(((candidates["stored_quote"] < 0) & (candidates["stored_quote"] != -100)).sum()),
            ),
            ("yas_price_override_requests", int(candidates["yas_request_eligible"].sum())),
            ("source_quote_matrix", str(CONFIG["YA_PATH"])),
            ("source_old_spread_output", str(old_spread_path)),
            ("source_submitted_panel", str(panel_path)),
            (
                "important_note",
                "stored_quote is a source-specific legacy PX_LAST and may "
                "represent either price/PU or yield depending on quote "
                "convention; small negative values are retained because "
                "the audit found continuous negative yield series, while "
                "zero and -100 are explicitly flagged for review. "
                "Classification is preliminary until Bloomberg YAS "
                "validation is completed.",
            ),
        ],
        columns=["metric", "value"],
    )

    instructions = pd.DataFrame(
        [
            (
                1,
                "Open this workbook in Excel with the Bloomberg Add-In active.",
            ),
            (
                2,
                "Open validation_sample and allow BDP formulas in "
                "yas_yield_if_treated_as_price to refresh only where "
                "yas_request_eligible is TRUE.",
            ),
            (
                3,
                "Wait until no cell shows '#N/A Requesting Data...'. Do not "
                "copy results while requests are still pending.",
            ),
            (
                4,
                "Review the ANCHOR rows first: EJ422135 Corp 2012-12-28 and "
                "AM250786 Corp 2017-06-30.",
            ),
            (
                5,
                "Use automated_check only as a diagnostic. Record the final "
                "decision in analyst_review: PRICE, YIELD, AMBIGUOUS, or ERROR.",
            ),
            (
                6,
                "After refreshing, save a separate copy and paste the "
                "validation_sample sheet as values only. Do not overwrite "
                "ya.v1.xlsx or the old spread output.",
            ),
            (
                7,
                "Return the values-only validation results before building "
                "the complete 4,651-row conversion request.",
            ),
        ],
        columns=["step", "instruction"],
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(
        output_path,
        engine="openpyxl",
        date_format="yyyy-mm-dd",
        datetime_format="yyyy-mm-dd",
    ) as writer:
        instructions.to_excel(
            writer,
            sheet_name="instructions",
            index=False,
        )
        summary.to_excel(writer, sheet_name="summary", index=False)
        classification_summary.to_excel(
            writer,
            sheet_name="classification_summary",
            index=False,
        )
        sample.to_excel(
            writer,
            sheet_name="validation_sample",
            index=False,
        )
        candidate_universe.to_excel(
            writer,
            sheet_name="candidate_universe",
            index=False,
        )
        bond_summary.to_excel(
            writer,
            sheet_name="bond_summary",
            index=False,
        )

    _style_workbook(output_path)

    print("Corporate YAS validation workbook created successfully.")
    print(f"Output: {output_path}")
    print(f"Thesis bonds: {unique_bonds}")
    print(f"Eligible observations: {eligible_rows}")
    print(f"Old IPCA output references: {old_output_rows}")
    print(f"Submitted panel rows: {old_panel_rows}")
    print(f"Rows without old output reference: {no_old_output_rows}")
    print(f"Validation sample rows: {len(sample)}")

    return output_path


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Build a Bloomberg/YAS validation sample for the mixed corporate "
            "quote matrix."
        )
    )
    parser.add_argument(
        "--panel",
        type=Path,
        default=DEFAULT_PANEL_PATH,
        help="Path to panel_main_and_robustness.xlsx.",
    )
    parser.add_argument(
        "--old-spreads",
        type=Path,
        default=DEFAULT_OLD_SPREAD_PATH,
        help="Path to output_spreads_di_and_ipca.xlsx.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=DEFAULT_OUTPUT_PATH,
        help="Output validation workbook path.",
    )
    parser.add_argument(
        "--per-group",
        type=int,
        default=8,
        help="Number of diverse observations selected per validation group.",
    )
    parser.add_argument(
        "--no-strict",
        action="store_true",
        help=(
            "Allow audited row counts to differ. Validation invariants still "
            "fail. Use only after investigating the reason."
        ),
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    build_validation_workbook(
        panel_path=args.panel.resolve(),
        old_spread_path=args.old_spreads.resolve(),
        output_path=args.output.resolve(),
        per_group=args.per_group,
        strict=not args.no_strict,
    )


if __name__ == "__main__":
    main()
