#!/usr/bin/env python
"""Build the selected pricing-source provenance table for the 103 thesis bonds.

This script reconstructs the historical pricing-source selection from the
20 preserved waterfall result workbooks. It reads only the header row of each
workbook; it does not load or modify the historical quote series.

Expected audited snapshot
-------------------------
- 20 waterfall result files
- 4,198 header occurrences across all files
- 3,952 unique corporate securities
- 246 duplicated securities (the overlap is expected; notably parts 3/4)
- 0 conflicting selected sources among duplicates
- 103 thesis bonds, all matched to a historical selected source

Output
------
datos_y_modelos/db/brazil_domestic_corp_bonds/
    corp_selected_pricing_sources.v1.xlsx
"""

from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


REPO_ROOT = Path(__file__).resolve().parents[1]

DEFAULT_WATERFALL_DIR = (
    REPO_ROOT
    / "datos_y_modelos"
    / "db"
    / "brazil_domestic_corp_bonds"
    / "pricing_source_waterfall_results"
)

DEFAULT_PANEL_PATH = (
    REPO_ROOT
    / "datos_y_modelos"
    / "db"
    / "econometrics_db"
    / "panel_main_and_robustness.xlsx"
)

DEFAULT_OUTPUT_PATH = (
    REPO_ROOT
    / "datos_y_modelos"
    / "db"
    / "brazil_domestic_corp_bonds"
    / "corp_selected_pricing_sources.v1.xlsx"
)

EXPECTED_FILES = 20
EXPECTED_HEADER_OCCURRENCES = 4_198
EXPECTED_UNIQUE_BONDS = 3_952
EXPECTED_DUPLICATED_BONDS = 246
EXPECTED_THESIS_BONDS = 103
EXPECTED_SOURCE_COUNTS = {
    "B3CY": 41,
    "ANDE": 28,
    "ATIV": 15,
    "BRAD": 13,
    "PSAN": 4,
    "BVAL": 1,
    "CMDB": 1,
}


class ProvenanceError(RuntimeError):
    """Raised when the preserved waterfall files fail an audit check."""


def _clean(value: object) -> str:
    return str(value).strip()


def _part_number(path: Path) -> int:
    stem = path.stem
    try:
        return int(stem.rsplit("_", 1)[-1])
    except ValueError as exc:
        raise ProvenanceError(
            f"Unexpected waterfall filename: {path.name}"
        ) from exc


def _find_waterfall_files(folder: Path) -> list[Path]:
    files = sorted(
        folder.glob("resultado_parte_*.xlsx"),
        key=_part_number,
    )

    if not files:
        raise FileNotFoundError(
            f"No resultado_parte_*.xlsx files found in: {folder}"
        )

    part_numbers = [_part_number(path) for path in files]
    duplicates = [n for n, c in Counter(part_numbers).items() if c > 1]
    missing = sorted(set(range(1, EXPECTED_FILES + 1)) - set(part_numbers))

    if duplicates:
        raise ProvenanceError(
            f"Duplicate waterfall part numbers found: {duplicates}"
        )
    if missing:
        raise ProvenanceError(
            f"Missing waterfall result parts: {missing}"
        )

    return files


def _read_headers(path: Path) -> list[str]:
    """Read non-empty first-row headers without loading the large data matrix."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[workbook.sheetnames[0]]
        first_row = next(
            worksheet.iter_rows(min_row=1, max_row=1, values_only=True)
        )
        return [
            _clean(value)
            for value in first_row
            if value is not None and _clean(value)
        ]
    finally:
        workbook.close()


def _parse_header(header: str, source_file: str) -> dict[str, str]:
    if header.upper() == "DATE":
        raise ValueError("DATE is not a security header.")
    if "@" not in header:
        raise ProvenanceError(
            f"Header has no @pricing-source suffix: {header!r} "
            f"in {source_file}"
        )

    base_id, selected_source = header.rsplit("@", 1)
    base_id = base_id.strip()
    selected_source = selected_source.strip().upper()

    if not base_id or not selected_source:
        raise ProvenanceError(
            f"Malformed security header: {header!r} in {source_file}"
        )

    return {
        "base_id": base_id,
        "selected_source": selected_source,
        "source_file": source_file,
        "original_header": header,
    }


def _style_workbook(path: Path) -> None:
    """Apply light audit-friendly formatting to the generated workbook."""
    wb = load_workbook(path)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.sheet_view.showGridLines = False

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        widths: dict[int, int] = {}
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                widths[cell.column] = min(
                    max(widths.get(cell.column, 0), len(str(cell.value)) + 2),
                    55,
                )

        for col_idx, width in widths.items():
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        ws.auto_filter.ref = ws.dimensions

    wb.save(path)


def build_selected_source_table(
    waterfall_dir: Path,
    panel_path: Path,
    output_path: Path,
    strict: bool = True,
) -> Path:
    if not waterfall_dir.exists():
        raise FileNotFoundError(
            f"Waterfall result folder not found: {waterfall_dir}"
        )
    if not panel_path.exists():
        raise FileNotFoundError(f"Econometric panel not found: {panel_path}")

    files = _find_waterfall_files(waterfall_dir)

    # ------------------------------------------------------------------
    # 1. Recover every selected source from workbook headers.
    # ------------------------------------------------------------------
    header_records: list[dict[str, str]] = []

    for path in files:
        for header in _read_headers(path):
            if header.upper() == "DATE":
                continue
            header_records.append(_parse_header(header, path.name))

    headers = pd.DataFrame(header_records)
    if headers.empty:
        raise ProvenanceError("No security headers were recovered.")

    grouped: dict[str, list[dict[str, str]]] = defaultdict(list)
    for record in header_records:
        grouped[record["base_id"]].append(record)

    conflicting: list[dict[str, str]] = []
    dedup_rows: list[dict[str, object]] = []
    duplicate_audit_rows: list[dict[str, object]] = []

    for base_id in sorted(grouped):
        records = grouped[base_id]
        sources = sorted({record["selected_source"] for record in records})
        files_seen = sorted({record["source_file"] for record in records})
        headers_seen = sorted({record["original_header"] for record in records})

        if len(sources) != 1:
            conflicting.append(
                {
                    "base_id": base_id,
                    "sources": "; ".join(sources),
                    "files": "; ".join(files_seen),
                }
            )
            continue

        dedup_rows.append(
            {
                "base_id": base_id,
                "selected_source": sources[0],
                "security_with_source": f"{base_id}@{sources[0]} Corp",
                "header_occurrences": len(records),
                "duplicate_in_waterfall": len(records) > 1,
                "source_files": "; ".join(files_seen),
                "original_headers": "; ".join(headers_seen),
            }
        )

        if len(records) > 1:
            duplicate_audit_rows.append(
                {
                    "base_id": base_id,
                    "selected_source": sources[0],
                    "header_occurrences": len(records),
                    "source_files": "; ".join(files_seen),
                    "source_consistent": True,
                }
            )

    if conflicting:
        conflict_df = pd.DataFrame(conflicting)
        raise ProvenanceError(
            "Conflicting selected sources found for duplicated securities. "
            f"Examples:\n{conflict_df.head(20).to_string(index=False)}"
        )

    all_sources = pd.DataFrame(dedup_rows)
    duplicate_audit = pd.DataFrame(duplicate_audit_rows)

    # ------------------------------------------------------------------
    # 2. Restrict to the 103 bonds actually used in the submitted thesis.
    # ------------------------------------------------------------------
    panel = pd.read_excel(panel_path, sheet_name="ts", usecols=["bond_id"])
    panel["bond_id"] = panel["bond_id"].map(_clean)
    thesis_ids = sorted(panel["bond_id"].dropna().unique().tolist())

    panel_map = pd.DataFrame({"bond_id": thesis_ids})
    panel_map["base_id"] = panel_map["bond_id"].str.replace(
        r"\s+Corp$", "", regex=True
    )

    selected = panel_map.merge(
        all_sources,
        on="base_id",
        how="left",
        validate="one_to_one",
    )

    missing = selected[selected["selected_source"].isna()]["bond_id"].tolist()
    if missing:
        raise ProvenanceError(
            "Thesis bonds missing from preserved waterfall results: "
            f"{missing}"
        )

    selected = selected[
        [
            "bond_id",
            "base_id",
            "selected_source",
            "security_with_source",
            "header_occurrences",
            "duplicate_in_waterfall",
            "source_files",
            "original_headers",
        ]
    ].sort_values("bond_id").reset_index(drop=True)

    # ------------------------------------------------------------------
    # 3. Audited counts and strict snapshot checks.
    # ------------------------------------------------------------------
    actual_source_counts = (
        selected["selected_source"].value_counts().sort_index().to_dict()
    )

    actual_counts = {
        "waterfall_files": len(files),
        "header_occurrences": len(headers),
        "unique_waterfall_bonds": len(all_sources),
        "duplicated_waterfall_bonds": int(
            all_sources["duplicate_in_waterfall"].sum()
        ),
        "conflicting_selected_sources": 0,
        "thesis_bonds": len(selected),
        "thesis_bonds_missing_source": len(missing),
    }

    expected_counts = {
        "waterfall_files": EXPECTED_FILES,
        "header_occurrences": EXPECTED_HEADER_OCCURRENCES,
        "unique_waterfall_bonds": EXPECTED_UNIQUE_BONDS,
        "duplicated_waterfall_bonds": EXPECTED_DUPLICATED_BONDS,
        "conflicting_selected_sources": 0,
        "thesis_bonds": EXPECTED_THESIS_BONDS,
        "thesis_bonds_missing_source": 0,
    }

    count_mismatches = {
        key: (actual_counts[key], expected_counts[key])
        for key in expected_counts
        if actual_counts[key] != expected_counts[key]
    }

    source_mismatches = {
        source: (
            actual_source_counts.get(source, 0),
            expected_count,
        )
        for source, expected_count in EXPECTED_SOURCE_COUNTS.items()
        if actual_source_counts.get(source, 0) != expected_count
    }

    unexpected_sources = sorted(
        set(actual_source_counts) - set(EXPECTED_SOURCE_COUNTS)
    )

    if strict and (count_mismatches or source_mismatches or unexpected_sources):
        raise ProvenanceError(
            "Recovered provenance differs from the audited snapshot. "
            f"Count mismatches={count_mismatches}; "
            f"source mismatches={source_mismatches}; "
            f"unexpected sources={unexpected_sources}. "
            "Investigate before using --no-strict."
        )

    source_summary = (
        selected.groupby("selected_source", as_index=False)
        .agg(thesis_bonds=("bond_id", "nunique"))
        .sort_values("thesis_bonds", ascending=False)
        .reset_index(drop=True)
    )
    source_summary["share_pct"] = (
        source_summary["thesis_bonds"] / len(selected) * 100
    )

    summary_rows = [
        ("waterfall_folder", str(waterfall_dir.relative_to(REPO_ROOT))),
        ("source_panel", str(panel_path.relative_to(REPO_ROOT))),
        ("waterfall_files", len(files)),
        ("header_occurrences", len(headers)),
        ("unique_waterfall_bonds", len(all_sources)),
        (
            "duplicated_waterfall_bonds",
            int(all_sources["duplicate_in_waterfall"].sum()),
        ),
        ("conflicting_selected_sources", 0),
        ("thesis_bonds", len(selected)),
        ("thesis_bonds_missing_source", len(missing)),
        (
            "method",
            "Selected pricing source recovered from the @SOURCE suffix "
            "preserved in each waterfall result workbook header.",
        ),
        (
            "next_step",
            "Use security_with_source to retrieve a homogeneous "
            "YLD_YTM_MID series in BQuant, retaining quote date and "
            "staleness_days.",
        ),
    ]
    summary = pd.DataFrame(summary_rows, columns=["metric", "value"])

    # ------------------------------------------------------------------
    # 4. Save workbook.
    # ------------------------------------------------------------------
    output_path.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="summary", index=False)
        selected.to_excel(writer, sheet_name="selected_sources", index=False)
        source_summary.to_excel(writer, sheet_name="source_summary", index=False)
        duplicate_audit.to_excel(
            writer,
            sheet_name="duplicate_audit",
            index=False,
        )
        all_sources.to_excel(
            writer,
            sheet_name="all_waterfall_sources",
            index=False,
        )

    _style_workbook(output_path)

    print("Corporate selected pricing-source table created successfully.")
    print(f"Output: {output_path}")
    print(f"Waterfall files: {len(files)}")
    print(f"Header occurrences: {len(headers)}")
    print(f"Unique waterfall bonds: {len(all_sources)}")
    print(
        "Duplicated waterfall bonds: "
        f"{int(all_sources['duplicate_in_waterfall'].sum())}"
    )
    print("Conflicting selected sources: 0")
    print(f"Thesis bonds matched: {len(selected)}")
    print("\nThesis source distribution:")
    print(source_summary.to_string(index=False))

    return output_path


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Recover the historical selected pricing source for each of the "
            "103 thesis corporate bonds."
        )
    )
    parser.add_argument(
        "--waterfall-dir",
        type=Path,
        default=DEFAULT_WATERFALL_DIR,
        help="Folder containing resultado_parte_1.xlsx ... resultado_parte_20.xlsx.",
    )
    parser.add_argument(
        "--panel",
        type=Path,
        default=DEFAULT_PANEL_PATH,
        help="Path to panel_main_and_robustness.xlsx.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=DEFAULT_OUTPUT_PATH,
        help="Output XLSX path.",
    )
    parser.add_argument(
        "--no-strict",
        action="store_true",
        help=(
            "Allow audited counts to differ. Structural validation errors "
            "still fail."
        ),
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    build_selected_source_table(
        waterfall_dir=args.waterfall_dir.resolve(),
        panel_path=args.panel.resolve(),
        output_path=args.output.resolve(),
        strict=not args.no_strict,
    )


if __name__ == "__main__":
    main()
