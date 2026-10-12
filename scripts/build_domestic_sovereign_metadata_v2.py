from pathlib import Path
from copy import copy

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]

OLD = ROOT / "datos_y_modelos" / "Domestic" / "domestic_sovereign_curve_brazil.xlsx"
HIST = ROOT / "datos_y_modelos" / "db" / "brazil_domestic_govt_bonds" / "historical_ltn_ntnf_2010_2015_yas_formulas.xlsx"
OUT = ROOT / "datos_y_modelos" / "Domestic" / "domestic_sovereign_curve_brazil.v2.xlsx"

SHEET = "db_values_only"

# ----------------------------------------------------------------------
# Read and validate existing metadata
# ----------------------------------------------------------------------
old = pd.read_excel(OLD, sheet_name=SHEET)
hist = pd.read_excel(HIST, sheet_name="metadata_value")

assert len(old) == 46
assert len(hist) == 42

required_hist = [
    "bond_id",
    "id_bb_global",
    "ID_ISIN",
    "issue_dt",
    "maturity",
    "derived_type",
]

assert hist[required_hist].notna().all().all()

hist["bond_id"] = hist["bond_id"].astype(str).str.strip()
hist["id_bb_global"] = hist["id_bb_global"].astype(str).str.strip()
hist["ID_ISIN"] = hist["ID_ISIN"].astype(str).str.strip()
hist["derived_type"] = hist["derived_type"].astype(str).str.upper().str.strip()

assert hist["derived_type"].value_counts().to_dict() == {
    "LTN": 31,
    "NTNF": 11,
}

assert hist["ID_ISIN"].nunique() == 42
assert hist["id_bb_global"].nunique() == 42

existing_isins = set(old["ID_ISIN"].dropna().astype(str).str.strip())
assert not (existing_isins & set(hist["ID_ISIN"]))

# ----------------------------------------------------------------------
# Canonical identifier used by govt_ya_ande.v3
# ----------------------------------------------------------------------
hist["canonical_id"] = hist["ID_ISIN"].apply(
    lambda x: x if x.endswith(" Corp") else x + " Corp"
)

assert hist["canonical_id"].nunique() == 42

# ----------------------------------------------------------------------
# Open original workbook and append rows
# ----------------------------------------------------------------------
wb = load_workbook(OLD)
ws = wb[SHEET]

nonempty_rows = [
    r for r in range(1, ws.max_row + 1)
    if any(ws.cell(r, c).value is not None for c in range(1, ws.max_column + 1))
]
last_data_row = max(nonempty_rows)

assert len(nonempty_rows) == 47
assert last_data_row == 47
assert ws.max_column == len(old.columns)

headers = {
    ws.cell(row=1, column=c).value: c
    for c in range(1, ws.max_column + 1)
}

required_output = [
    "ID",
    "id",
    "ISSUER",
    "ID_BB_GLOBAL",
    "ID_ISIN",
    "ISSUE_DT",
    "MATURITY",
    "CRNCY",
    "CPN_TYP",
    "CPN",
    "INFLATION_LINKED_INDICATOR",
    "CALC_TYP",
    "CALC_TYP_DES",
]

for col in required_output:
    assert col in headers

# Use an existing LTN row only as a formatting template.
template_row = None
for r in range(2, last_data_row + 1):
    if str(ws.cell(r, headers["CALC_TYP_DES"]).value).strip().upper() == "BRAZIL: BBCS/LTNS":
        template_row = r
        break

assert template_row is not None

start_row = last_data_row + 1

for i, rec in hist.reset_index(drop=True).iterrows():
    r = start_row + i

    # Copy formatting only, never existing values.
    for c in range(1, ws.max_column + 1):
        source = ws.cell(template_row, c)
        target = ws.cell(r, c)

        if source.has_style:
            target._style = copy(source._style)

        target.number_format = source.number_format
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)

    bond_type = rec["derived_type"]

    if bond_type == "LTN":
        issuer = "LETRA TESOURO NACIONAL"
        cpn_typ = "ZERO COUPON"
        cpn = 0
        calc_typ = 712
        calc_typ_des = "BRAZIL: BBCs/LTNs"

    elif bond_type == "NTNF":
        issuer = "NOTA DO TESOURO NACIONAL"
        cpn_typ = "FIXED"
        cpn = 10
        calc_typ = 1309
        calc_typ_des = "BRAZIL FIXED CPN"

    else:
        raise AssertionError(f"Unexpected bond type: {bond_type}")

    values = {
        "ID": rec["canonical_id"],
        "id": rec["canonical_id"],
        "ISSUER": issuer,
        "ID_BB_GLOBAL": rec["id_bb_global"],
        "ID_ISIN": rec["ID_ISIN"],
        "ISSUE_DT": pd.Timestamp(rec["issue_dt"]).to_pydatetime(),
        "MATURITY": pd.Timestamp(rec["maturity"]).to_pydatetime(),
        "CRNCY": "BRL",
        "CPN_TYP": cpn_typ,
        "CPN": cpn,
        "INFLATION_LINKED_INDICATOR": "N",
        "CALC_TYP": calc_typ,
        "CALC_TYP_DES": calc_typ_des,
    }

    for name, value in values.items():
        ws.cell(r, headers[name], value=value)

wb.save(OUT)

# ----------------------------------------------------------------------
# Re-open and validate final metadata
# ----------------------------------------------------------------------
check = pd.read_excel(OUT, sheet_name=SHEET)

assert len(check) == 88
old_check = old.reset_index(drop=True)
new_check = check.iloc[:46].reset_index(drop=True)

assert old_check.columns.tolist() == new_check.columns.tolist()

for col in old_check.columns:
    a = old_check[col]
    b = new_check[col]

    if col in {"ISSUE_DT", "MATURITY"}:
        aa = pd.to_datetime(a, errors="coerce")
        bb = pd.to_datetime(b, errors="coerce")
        assert aa.equals(bb), col

    elif pd.api.types.is_numeric_dtype(a) and pd.api.types.is_numeric_dtype(b):
        pd.testing.assert_series_equal(
            a.astype(float),
            b.astype(float),
            check_dtype=False,
            check_names=False,
            check_exact=False,
            rtol=0,
            atol=1e-9,
        )

    else:
        aa = a.fillna("<NA>").astype(str).str.strip()
        bb = b.fillna("<NA>").astype(str).str.strip()
        assert aa.equals(bb), col

appended = check.iloc[46:].copy()

assert len(appended) == 42
assert appended["id"].nunique() == 42
assert set(appended["id"]) == set(hist["canonical_id"])
assert appended["ID_ISIN"].nunique() == 42
assert appended["ID_BB_GLOBAL"].nunique() == 42

assert (
    appended["CALC_TYP_DES"]
    .value_counts()
    .to_dict()
    == {
        "BRAZIL: BBCs/LTNs": 31,
        "BRAZIL FIXED CPN": 11,
    }
)

assert (
    appended["INFLATION_LINKED_INDICATOR"]
    .astype(str)
    .str.upper()
    .eq("N")
    .all()
)

print("Created extended sovereign metadata workbook.")
print("Output:", OUT)
print("Existing metadata rows preserved:", 46)
print("Historical rows appended:", 42)
print("  LTN:", 31)
print("  NTN-F:", 11)
print("Final metadata rows:", len(check))
print("Unique appended ISINs:", appended["ID_ISIN"].nunique())
print("Unique appended BBGIDs:", appended["ID_BB_GLOBAL"].nunique())
