from pathlib import Path
from copy import copy
import os
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]

HIST = ROOT / r"datos_y_modelos/db/brazil_domestic_govt_bonds/historical_ltn_ntnf_2010_2015_yas_formulas.xlsx"
TARGET = ROOT / r"datos_y_modelos/db/brazil_domestic_govt_bonds/govt_ya_ande.v3.xlsx"
TEMP = TARGET.with_name("govt_ya_ande.v3.tmp.xlsx")


def norm_date(value):
    if value is None:
        return None
    if hasattr(value, "date"):
        return value.date()
    raise TypeError(f"Unexpected date value: {value!r}")


def norm_id(value):
    return str(value).strip() if value is not None else ""


# ------------------------------------------------------------
# Read historical workbook
# ------------------------------------------------------------

wh = load_workbook(HIST, read_only=True, data_only=True)

assert "metadata_value" in wh.sheetnames
assert "ya_value" in wh.sheetnames

wm = wh["metadata_value"]
wy = wh["ya_value"]

# Metadata headers
meta_headers = {
    norm_id(wm.cell(1, c).value): c
    for c in range(1, wm.max_column + 1)
}

assert "bond_id" in meta_headers
assert "ID_ISIN" in meta_headers

bond_col = meta_headers["bond_id"]
isin_col = meta_headers["ID_ISIN"]

# bond_id -> ISIN Corp
bond_to_isin = {}

for r in range(2, wm.max_row + 1):
    bond_id = norm_id(wm.cell(r, bond_col).value)
    isin = norm_id(wm.cell(r, isin_col).value)

    if not bond_id or not isin:
        continue

    isin_header = isin if isin.endswith(" Corp") else isin + " Corp"

    if bond_id in bond_to_isin:
        raise AssertionError(f"Duplicate bond_id in metadata: {bond_id}")

    bond_to_isin[bond_id] = isin_header

assert len(bond_to_isin) == 42, (
    f"Expected 42 historical bond mappings, got {len(bond_to_isin)}"
)
assert len(set(bond_to_isin.values())) == 42


# ------------------------------------------------------------
# Read historical YA matrix
# Row 4 = headers
# Column C = observation date
# Columns D onward = historical bond IDs
# ------------------------------------------------------------

hist_bonds = []

for c in range(4, wy.max_column + 1):
    bond_id = norm_id(wy.cell(4, c).value)

    if not bond_id:
        continue

    if bond_id not in bond_to_isin:
        raise AssertionError(
            f"Historical YA bond missing from metadata_value: {bond_id}"
        )

    hist_bonds.append(
        (c, bond_id, bond_to_isin[bond_id])
    )

assert len(hist_bonds) == 42, (
    f"Expected 42 historical YA columns, got {len(hist_bonds)}"
)

hist_dates = {}

for r in range(5, wy.max_row + 1):
    d = norm_date(wy.cell(r, 3).value)

    if d is None:
        continue

    if d in hist_dates:
        raise AssertionError(f"Duplicate historical date: {d}")

    hist_dates[d] = r

assert len(hist_dates) == 61, (
    f"Expected 61 historical dates, got {len(hist_dates)}"
)


# ------------------------------------------------------------
# Open govt_ya_ande.v3 and preserve original block
# ------------------------------------------------------------

wt = load_workbook(TARGET)
assert "ya_values_only" in wt.sheetnames

ws = wt["ya_values_only"]

original_max_col = ws.max_column
original_max_row = ws.max_row

assert original_max_col == 32, (
    f"Expected original v3 to have 32 columns, got {original_max_col}"
)
assert original_max_row == 187, (
    f"Expected original v3 to have 187 rows, got {original_max_row}"
)

original_values = [
    [ws.cell(r, c).value for c in range(1, original_max_col + 1)]
    for r in range(1, original_max_row + 1)
]

# Existing headers
header_to_col = {}

for c in range(1, ws.max_column + 1):
    h = norm_id(ws.cell(1, c).value)
    if h:
        header_to_col[h] = c

# Existing dates
date_to_row = {}

for r in range(2, ws.max_row + 1):
    d = norm_date(ws.cell(r, 1).value)

    if d is None:
        continue

    if d in date_to_row:
        raise AssertionError(f"Duplicate target date: {d}")

    date_to_row[d] = r

missing_dates = sorted(set(hist_dates) - set(date_to_row))

assert not missing_dates, (
    "Historical dates absent from govt_ya_ande.v3: "
    + ", ".join(str(x) for x in missing_dates)
)


# ------------------------------------------------------------
# Append / identify the 42 ISIN columns
# ------------------------------------------------------------

template_col = original_max_col
template_letter = get_column_letter(template_col)

new_headers = 0
destination_cols = {}

for _, _, isin_header in hist_bonds:

    if isin_header in header_to_col:
        destination_cols[isin_header] = header_to_col[isin_header]
        continue

    new_col = ws.max_column + 1
    ws.cell(1, new_col).value = isin_header

    # Copy formatting convention from final existing column
    for r in range(1, ws.max_row + 1):
        src = ws.cell(r, template_col)
        dst = ws.cell(r, new_col)

        if src.has_style:
            dst._style = copy(src._style)

        dst.font = copy(src.font)
        dst.fill = copy(src.fill)
        dst.border = copy(src.border)
        dst.alignment = copy(src.alignment)
        dst.number_format = src.number_format
        dst.protection = copy(src.protection)

    new_letter = get_column_letter(new_col)
    ws.column_dimensions[new_letter].width = (
        ws.column_dimensions[template_letter].width
    )

    header_to_col[isin_header] = new_col
    destination_cols[isin_header] = new_col
    new_headers += 1


# ------------------------------------------------------------
# Transfer historical observations by exact date
# ------------------------------------------------------------

written = 0
nonzero = 0
zeros = 0

for source_col, bond_id, isin_header in hist_bonds:

    dest_col = destination_cols[isin_header]

    for d, source_row in hist_dates.items():

        dest_row = date_to_row[d]
        source_value = wy.cell(source_row, source_col).value
        dest_cell = ws.cell(dest_row, dest_col)

        # Normalize blank
        if source_value is None:
            value = None
        else:
            value = float(source_value)

        # Idempotency / conflict protection
        if dest_cell.value not in (None, ""):
            existing = float(dest_cell.value)

            if value is None or abs(existing - value) > 1e-10:
                raise AssertionError(
                    f"Conflict for {isin_header} on {d}: "
                    f"target={dest_cell.value}, source={source_value}"
                )

        else:
            dest_cell.value = value

        written += 1

        if value is not None and value != 0:
            nonzero += 1
        else:
            zeros += 1


# ------------------------------------------------------------
# Pre-save validations
# ------------------------------------------------------------

assert written == 61 * 42, (
    f"Expected 2562 transferred cells, got {written}"
)

assert nonzero == 1042, (
    f"Expected 1042 non-zero historical YAS values, got {nonzero}"
)

assert zeros == 1520, (
    f"Expected 1520 zero/blank historical cells, got {zeros}"
)

assert ws.max_column == 74, (
    f"Expected 74 total columns after consolidation, got {ws.max_column}"
)


# ------------------------------------------------------------
# Save to temporary file first
# ------------------------------------------------------------

if TEMP.exists():
    TEMP.unlink()

wt.save(TEMP)


# ------------------------------------------------------------
# Post-save audit
# ------------------------------------------------------------

wc = load_workbook(TEMP, read_only=True, data_only=True)
check = wc["ya_values_only"]

assert check.max_row == 187
assert check.max_column == 74

# Original v3 block must remain identical
for r in range(1, original_max_row + 1):
    for c in range(1, original_max_col + 1):
        if check.cell(r, c).value != original_values[r - 1][c - 1]:
            raise AssertionError(
                f"Original v3 cell changed unexpectedly: "
                f"{get_column_letter(c)}{r}"
            )

# Close read-only workbook before replacing file on Windows
wc.close()

# Atomic replacement only after every validation succeeds
os.replace(TEMP, TARGET)

print("Historical sovereign YAS consolidated into govt_ya_ande.v3.xlsx.")
print("Historical bonds mapped:", len(hist_bonds))
print("Historical dates matched:", len(hist_dates))
print("New ISIN columns added:", new_headers)
print("Cells transferred:", written)
print("Non-zero historical YAS:", nonzero)
print("Zero/blank historical cells:", zeros)
print("Final total columns:", 74)
print("Original v3 block preserved: YES")
print("Target:", TARGET)
