from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]

BONDS_PATH = (
    ROOT
    / "datos_y_modelos"
    / "db"
    / "brazil_domestic_govt_bonds"
    / "historical_ltn_ntnf_2010_2015.xlsx"
)

PANEL_PATH = (
    ROOT
    / "datos_y_modelos"
    / "db"
    / "econometrics_db"
    / "panel_main_and_robustness_rebuilt.xlsx"
)

OUTPUT_PATH = (
    ROOT
    / "datos_y_modelos"
    / "db"
    / "brazil_domestic_govt_bonds"
    / "historical_ltn_ntnf_2010_2015_yas_formulas.xlsx"
)


def main():
    bonds = pd.read_excel(BONDS_PATH)

    required = {
        "ID",
        "id_bb_global",
        "ticker",
        "issue_dt",
        "maturity",
    }
    missing = required.difference(bonds.columns)
    if missing:
        raise AssertionError(
            f"Historical sovereign universe missing columns: {sorted(missing)}"
        )

    bonds = bonds.copy()
    bonds["ID"] = bonds["ID"].astype(str).str.strip()
    bonds["issue_dt"] = pd.to_datetime(bonds["issue_dt"], errors="raise")
    bonds["maturity"] = pd.to_datetime(bonds["maturity"], errors="raise")

    if len(bonds) != 42:
        raise AssertionError(
            f"Expected 42 historical LTN/NTN-F securities, got {len(bonds)}."
        )

    if bonds["ID"].duplicated().any():
        raise AssertionError("Duplicate security IDs in historical universe.")

    panel = pd.read_excel(
        PANEL_PATH,
        sheet_name="ts",
        usecols=["obs_date", "sovereign_spread_brl"],
    )

    panel["obs_date"] = pd.to_datetime(
        panel["obs_date"],
        errors="raise",
    ).dt.normalize()

    dates = (
        panel.loc[
            panel["sovereign_spread_brl"].isna()
            & (panel["obs_date"] <= pd.Timestamp("2015-12-31")),
            "obs_date",
        ]
        .drop_duplicates()
        .sort_values()
        .tolist()
    )

    if len(dates) != 61:
        raise AssertionError(
            f"Expected 61 missing sovereign dates, got {len(dates)}."
        )

    if dates[0] != pd.Timestamp("2010-12-30"):
        raise AssertionError(
            f"Unexpected first date: {dates[0]}"
        )

    if dates[-1] != pd.Timestamp("2015-12-30"):
        raise AssertionError(
            f"Unexpected last date: {dates[-1]}"
        )

    eligible_pairs = 0
    for obs_date in dates:
        eligible_pairs += int(
            (
                (bonds["issue_dt"] <= obs_date)
                & (bonds["maturity"] > obs_date)
            ).sum()
        )

    if eligible_pairs != 1056:
        raise AssertionError(
            "Historical eligibility count changed: "
            f"expected 1056, got {eligible_pairs}."
        )

    wb = Workbook()
    wb.remove(wb.active)

    meta_ws = wb.create_sheet("metadata")
    hp_ws = wb.create_sheet("hp")
    ya_ws = wb.create_sheet("ya")

    # ----------------------------------------------------------
    # Metadata / audit sheet
    # ----------------------------------------------------------
    meta_headers = [
        "bond_id",
        "id_bb_global",
        "ticker",
        "issue_dt",
        "maturity",
        "derived_type",
    ]

    for col, value in enumerate(meta_headers, start=1):
        meta_ws.cell(1, col, value)

    for r, row in enumerate(bonds.itertuples(index=False), start=2):
        meta_ws.cell(r, 1, row.ID)
        meta_ws.cell(r, 2, row.id_bb_global)
        meta_ws.cell(r, 3, row.ticker)
        meta_ws.cell(r, 4, row.issue_dt.to_pydatetime())
        meta_ws.cell(r, 5, row.maturity.to_pydatetime())
        meta_ws.cell(
            r,
            6,
            "LTN" if row.ticker == "BLTN" else "NTNF",
        )

    meta_ws["H1"] = "source_convention"
    meta_ws["H2"] = "Historical price source"
    meta_ws["I2"] = "ANDE PX_LAST"
    meta_ws["H3"] = "Historical price date"
    meta_ws["I3"] = "OBS_DATE"
    meta_ws["H4"] = "YAS price override"
    meta_ws["I4"] = "YAS_BOND_PX = ANDE PX_LAST"
    meta_ws["H5"] = "YAS settlement date"
    meta_ws["I5"] = "SETTLE_DT = OBS_DATE"
    meta_ws["H6"] = "Eligible bond-date cells"
    meta_ws["I6"] = eligible_pairs

    # ----------------------------------------------------------
    # Formula surfaces
    # Layout mirrors the audited legacy workbook:
    # row 4 = bond IDs
    # column C = observation dates
    # ----------------------------------------------------------
    for ws in (hp_ws, ya_ws):
        ws["C4"] = "id"

        for j, bond_id in enumerate(bonds["ID"], start=4):
            ws.cell(4, j, bond_id)

        for i, obs_date in enumerate(dates, start=5):
            ws.cell(i, 3, obs_date.to_pydatetime())
            ws.cell(i, 3).number_format = "yyyy-mm-dd"

    hp_formula_count = 0
    ya_formula_count = 0

    for i, obs_date in enumerate(dates, start=5):
        for j, bond in enumerate(
            bonds.itertuples(index=False),
            start=4,
        ):
            active = (
                bond.issue_dt <= obs_date
                and bond.maturity > obs_date
            )

            hp_cell = hp_ws.cell(i, j)
            ya_cell = ya_ws.cell(i, j)

            if not active:
                hp_cell.value = 0
                ya_cell.value = 0
                continue

            col = get_column_letter(j)

            hp_cell.value = (
                f'=BDH('
                f'SUBSTITUTE({col}$4," Corp","@ANDE Corp"),'
                f'"PX_LAST",$C{i},$C{i},'
                f'"days=all","dts=h","fill=0","QtTyp=P")'
            )

            ya_cell.value = (
                f'=IF(hp!{col}{i}=0,0,'
                f'BDP({col}$4,"YAS_BOND_YLD",'
                f'"YAS_BOND_PX",hp!{col}{i},'
                f'"SETTLE_DT",TEXT($C{i},"YYYYMMDD")))'
            )

            hp_formula_count += 1
            ya_formula_count += 1

    if hp_formula_count != 1056:
        raise AssertionError(
            f"Expected 1056 HP formulas, got {hp_formula_count}."
        )

    if ya_formula_count != 1056:
        raise AssertionError(
            f"Expected 1056 YAS formulas, got {ya_formula_count}."
        )

    # ----------------------------------------------------------
    # Minimal formatting
    # ----------------------------------------------------------
    header_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7",
    )

    for ws in (meta_ws, hp_ws, ya_ws):
        ws.freeze_panes = "D5" if ws.title != "metadata" else "A2"
        ws.sheet_view.showGridLines = False

    for ws in (hp_ws, ya_ws):
        for cell in ws[4]:
            if cell.column >= 3:
                cell.font = Font(bold=True)
                cell.fill = header_fill
                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                )

        ws.column_dimensions["C"].width = 13

        for j in range(4, 4 + len(bonds)):
            ws.column_dimensions[get_column_letter(j)].width = 21

    for cell in meta_ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill

    meta_ws.column_dimensions["A"].width = 22
    meta_ws.column_dimensions["B"].width = 18
    meta_ws.column_dimensions["C"].width = 12
    meta_ws.column_dimensions["D"].width = 13
    meta_ws.column_dimensions["E"].width = 13
    meta_ws.column_dimensions["F"].width = 14
    meta_ws.column_dimensions["H"].width = 28
    meta_ws.column_dimensions["I"].width = 32

    for row in range(2, 2 + len(bonds)):
        meta_ws.cell(row, 4).number_format = "yyyy-mm-dd"
        meta_ws.cell(row, 5).number_format = "yyyy-mm-dd"

    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"

    wb.save(OUTPUT_PATH)

    print("Builder created historical sovereign YAS workbook.")
    print("Output:", OUTPUT_PATH)
    print("Historical securities:", len(bonds))
    print("Observation dates:", len(dates))
    print("Eligible bond-date observations:", eligible_pairs)
    print("HP formulas:", hp_formula_count)
    print("YAS formulas:", ya_formula_count)


if __name__ == "__main__":
    main()
